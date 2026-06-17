"""DOCX equivalents of pdf/render.py renderers — for employers/ATS that want
.docx instead of .pdf. Same structured input."""

from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor


_HEADING_BLUE = RGBColor(0x1A, 0x30, 0x50)


def _set_default_font(doc: Document, name: str = "Helvetica", size_pt: int = 10) -> None:
    style = doc.styles["Normal"]
    style.font.name = name
    style.font.size = Pt(size_pt)


def _section_header(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = _HEADING_BLUE


def render_resume_docx(data: dict, out_path: Path) -> Path:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    _set_default_font(doc)

    # Header
    p = doc.add_paragraph()
    name_run = p.add_run(data.get("name", ""))
    name_run.bold = True
    name_run.font.size = Pt(20)
    if data.get("contact_line"):
        c = doc.add_paragraph(data["contact_line"])
        c.runs[0].font.size = Pt(10)

    if data.get("summary"):
        _section_header(doc, "Professional Summary")
        doc.add_paragraph(data["summary"])

    skills = data.get("skills") or []
    if skills:
        _section_header(doc, "Core Skills")
        for cat in skills:
            p = doc.add_paragraph()
            r = p.add_run(f"{cat.get('category', '')}: ")
            r.bold = True
            p.add_run(", ".join(cat.get("items") or []))

    exp = data.get("experience") or []
    if exp:
        _section_header(doc, "Experience")
        for role in exp:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            r = p.add_run(f"{role.get('company', '')} — {role.get('location', '')}")
            r.bold = True
            r.font.size = Pt(10.5)
            p2 = doc.add_paragraph()
            r2 = p2.add_run(f"{role.get('title', '')}  ·  {role.get('dates', '')}")
            r2.italic = True
            r2.font.size = Pt(10)
            for bullet in role.get("bullets") or []:
                doc.add_paragraph(bullet, style="List Bullet")

    edu = data.get("education") or []
    if edu:
        _section_header(doc, "Education & Certifications")
        for line in edu:
            doc.add_paragraph(line, style="List Bullet")

    doc.save(str(out_path))
    return out_path


def render_cover_letter_docx(data: dict, out_path: Path, sender_name: str, sender_contact: str) -> Path:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    _set_default_font(doc, size_pt=11)

    # Sender header
    p = doc.add_paragraph()
    r = p.add_run(sender_name)
    r.bold = True
    r.font.size = Pt(20)
    if sender_contact:
        c = doc.add_paragraph(sender_contact)
        c.runs[0].font.size = Pt(10)

    if data.get("date"):
        doc.add_paragraph(data["date"])

    # Recipient
    recipient_lines = [
        x for x in (data.get("recipient"), data.get("company"), data.get("company_address")) if x
    ]
    if recipient_lines:
        p = doc.add_paragraph()
        for i, line in enumerate(recipient_lines):
            if i:
                p.add_run().add_break()
            p.add_run(line)

    if data.get("greeting"):
        doc.add_paragraph(data["greeting"])

    for para in (data.get("body_paragraphs") or []):
        doc.add_paragraph(para)

    doc.add_paragraph(data.get("closing", "Sincerely,"))
    doc.add_paragraph()  # signature space
    doc.add_paragraph(data.get("signature_name", sender_name))

    doc.save(str(out_path))
    return out_path
